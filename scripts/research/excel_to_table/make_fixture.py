"""Deterministic Excel input fixtures (no ArcGIS/GDAL implementation code)."""
from pathlib import Path
from datetime import date, datetime, time
import openpyxl
from openpyxl.workbook.defined_name import DefinedName
import xlwt
root = Path(__file__).resolve().parents[3] / 'testdata' / 'geoproc'
w = openpyxl.Workbook()
s = w.active
s.title = 'Data'
for row in [
    ['small', 'large', 'decimal', 'label', 'mixed', 'flag', 'missing'],
    [1, 40000, 1.25, 'Alpha', 2, True, None],
    [-2, -50000, 3.0, 'βeta', 'three', False, '#N/A'],
    [32767, 2147483648, -2.5, 'A longer label', None, None, None],
]: s.append(row)
s = w.create_sheet('Names')
for row in [['bad name', 'bad-name', None, '1st', 'OBJECTID', 'select'],
            [10, 20, 30, 40, 50, 60], [11, 21, None, 41, 51, 61]]: s.append(row)
s = w.create_sheet('Offset')
for row in [[], ['preamble'], ['id', 'name'], [1, 'one'], [], [2, 'two']]: s.append(row)
s = w.create_sheet('Dates')
s.append(['day', 'clock', 'stamp'])
s.append([date(2024, 2, 29), time(13, 45, 30), datetime(2024, 2, 29, 13, 45, 30)])
s.append([None, time(0, 0), datetime(2025, 1, 2)])
s = w.create_sheet('Empty')
w.defined_names.add(DefinedName('Picked', attr_text="'Data'!$B$1:$D$3"))
w.save(root / 'excel_input.xlsx')
w = xlwt.Workbook()
s = w.add_sheet('Data')
for r, row in enumerate([['id', 'value', 'label'], [1, 1.25, 'one'], [-2, 3, 'two'], [None, None, 'three']]):
    for c, value in enumerate(row):
        if value is not None: s.write(r,c,value)
w.save(str(root / 'excel_input.xls'))
